"""Read PDF / image / TXT / MD / DOCX files into a list of (page_number, text).

Page numbers are real for PDFs. Formats without pages (TXT, MD, DOCX, images) are a single
"page 0", which the UI and citations treat as "no page".

PDFs and images (adapted from github.com/Aashrawat/CHATBOT):
  * PDF text comes from PyMuPDF (reading-order sorted, hyphenation repaired).
  * Pages with no text layer (scans) are rendered and transcribed by a GPT vision model.
  * Pages with little text but a large picture (slides, charts) get the picture described too.
  * Images are EXIF-rotated, downscaled, then described + OCR'd by the vision model.
"""

import base64
import io
import re
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from core.kb import config

MIN_TEXT_CHARS = 40  # below this a PDF page is treated as scanned
SHORT_PAGE_CHARS = 400  # pages this short with a big picture get the picture described too
VISION_WORKERS = 4


class EmptyDocument(Exception):
    """The file was readable but contains no extractable text (e.g. a blank PDF)."""


class TransientLoadError(Exception):
    """Reading failed for a reason that may go away (API key missing, rate limit, network).

    Unlike a corrupt file, this must not be remembered as a permanent failure.
    """


# --------------------------------------------------------------------------- vision
IMAGE_INDEX_PROMPT = """You are preparing an image for a search index. Produce exactly two sections.

DESCRIPTION: A thorough, factual description of what the image shows: objects, people, scene, layout, colours, \
and the meaning of any chart, diagram, screenshot or UI. Mention distinctive details someone might search for.

TEXT: Transcribe ALL visible text verbatim, in reading order. Render tables as Markdown tables. \
If there is no text, write "None".

Do not speculate beyond what is visible."""

PAGE_OCR_PROMPT = """Transcribe this document page faithfully and completely, in reading order.
- Keep headings, lists and paragraphs. Render tables as Markdown tables.
- Replace each figure, chart, diagram or photo with a one or two sentence description in the form \
[Figure: ...], including key numbers or trends for charts.
- Output only the page content, with no commentary."""


def _vision(prompt: str, jpeg: bytes) -> str:
    from core.kb.clients import sync_client

    b64 = base64.b64encode(jpeg).decode()
    response = sync_client().chat.completions.create(
        model=config.VISION_MODEL,
        messages=[{
            "role": "user",
            "content": [
                {"type": "text", "text": prompt},
                {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}", "detail": "high"}},
            ],
        }],
    )
    return (response.choices[0].message.content or "").strip()


# --------------------------------------------------------------------------- helpers
def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"-\n(?=[a-z])", "", text)  # de-hyphenate line wraps
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def prepare_image(data: bytes) -> bytes:
    """Decode any image, fix EXIF rotation, flatten transparency, downscale, return JPEG bytes."""
    from PIL import Image, ImageOps

    try:
        img = Image.open(io.BytesIO(data))
        img = ImageOps.exif_transpose(img)  # also loads the first frame of GIF / multi-page TIFF
    except Exception as exc:
        raise ValueError("That file is not a readable image.") from exc
    if img.mode in ("RGBA", "LA", "P"):
        rgba = img.convert("RGBA")
        background = Image.new("RGB", rgba.size, "white")
        background.paste(rgba, mask=rgba.split()[-1])
        img = background
    else:
        img = img.convert("RGB")
    img.thumbnail((config.MAX_IMAGE_SIDE, config.MAX_IMAGE_SIDE))
    out = io.BytesIO()
    img.save(out, "JPEG", quality=88, optimize=True)
    return out.getvalue()


def _vision_or_transient(prompt: str, jpeg: bytes) -> str:
    try:
        return _vision(prompt, jpeg)
    except Exception as exc:
        raise TransientLoadError(f"Vision model call failed: {type(exc).__name__}: {exc}") from exc


# --------------------------------------------------------------------------- PDF
def _has_big_image(page) -> bool:
    import pymupdf

    area = page.rect.width * page.rect.height
    if area <= 0:
        return False
    try:
        for info in page.get_image_info():
            rect = pymupdf.Rect(info["bbox"])
            if rect.width * rect.height > 0.25 * area:
                return True
    except Exception:
        pass
    return False


def _load_pdf(path: Path) -> list[tuple[int, str]]:
    try:
        import pymupdf
    except ImportError:  # keep working (text only) if PyMuPDF isn't installed
        return _load_pdf_text_only(path)

    try:
        doc = pymupdf.open(str(path))
    except Exception as exc:
        raise ValueError("PDF could not be opened (it may be corrupted)") from exc
    if doc.needs_pass:
        doc.close()
        raise ValueError("PDF is password protected")

    numbers: list[int] = []
    texts: list[str] = []
    modes: list[str | None] = []  # None | "replace" (scanned) | "augment" (text + figure)
    images: dict[int, bytes] = {}  # index -> rendered page, only for pages that need vision
    budget = config.MAX_VISION_PAGES
    try:
        for i, page in enumerate(doc):
            text = _clean(page.get_text("text", sort=True))
            mode = None
            if len(text) < MIN_TEXT_CHARS:
                mode = "replace"
            elif len(text) < SHORT_PAGE_CHARS and _has_big_image(page):
                mode = "augment"
            if mode and budget > 0:
                images[i] = page.get_pixmap(dpi=150).tobytes("jpeg")
                budget -= 1
            numbers.append(i + 1)
            texts.append(text)
            modes.append(mode)
    finally:
        doc.close()

    def read_visually(index: int) -> str | None:
        try:
            return _vision(PAGE_OCR_PROMPT, images[index])
        except Exception as exc:
            if modes[index] == "replace" and not texts[index]:
                # nothing else to fall back on, so don't mark the file failed
                raise TransientLoadError(
                    f"Page {numbers[index]} is scanned and the vision model call failed: {type(exc).__name__}: {exc}"
                ) from exc
            return None  # keep whatever text layer the page has

    if images:
        with ThreadPoolExecutor(max_workers=VISION_WORKERS) as pool:
            for index, vision in zip(images, pool.map(read_visually, images)):
                if modes[index] == "replace":
                    texts[index] = vision or texts[index]
                elif vision:
                    texts[index] = f"{texts[index]}\n\n[Visual content on this page]\n{vision}"

    return [(n, t) for n, t in zip(numbers, texts) if t]


def _load_pdf_text_only(path: Path) -> list[tuple[int, str]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ValueError("PDF is password protected") from exc
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append((number, text))
    return pages


# --------------------------------------------------------------------------- images
def _load_image(path: Path) -> list[tuple[int, str]]:
    jpeg = prepare_image(path.read_bytes())
    description = _vision_or_transient(IMAGE_INDEX_PROMPT, jpeg)
    # Prefix the filename so the image is findable by name and the answer can cite it
    return [(0, f"Image: {path.name}\n\n{description}")] if description else []


# --------------------------------------------------------------------------- text formats
def _load_text(path: Path) -> list[tuple[int, str]]:
    text = path.read_text(encoding="utf-8-sig", errors="replace").strip()
    return [(0, text)] if text else []


def _load_docx(path: Path) -> list[tuple[int, str]]:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    return [(0, text)] if text else []


_LOADERS = {
    ".pdf": _load_pdf,
    ".txt": _load_text,
    ".md": _load_text,
    ".docx": _load_docx,
    **{ext: _load_image for ext in config.IMAGE_EXTENSIONS},
}


def load_document(path: Path) -> list[tuple[int, str]]:
    loader = _LOADERS.get(path.suffix.lower())
    if loader is None:
        raise ValueError(f"Unsupported file type '{path.suffix}'")
    pages = loader(path)
    if not pages:
        raise EmptyDocument("No extractable text found")
    return pages
